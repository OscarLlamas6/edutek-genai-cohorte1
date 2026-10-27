"""Genera los documentos de ejemplo en formatos DOCX y PDF (los datos son ficticios).

Los archivos generados ya vienen en esta carpeta; este script solo sirve para regenerarlos o
para ver cómo se construyeron. Necesita `python-docx` y `reportlab` (grupo `dev`).

    uv run python datos/generar_documentos.py
"""

from pathlib import Path

from docx import Document
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

AQUI = Path(__file__).resolve().parent


def crear_docx() -> None:
    doc = Document()
    doc.add_heading("Condiciones generales del servicio de Andes Fibra", level=0)
    doc.add_paragraph(
        "Versión 2026-10. Documento ficticio creado para el curso. Describe las condiciones que se aplican según el "
        "tipo de cliente. Las cifras son inventadas."
    )

    doc.add_heading("Parte A · Clientes residenciales", level=1)
    doc.add_heading("Tiempo de reparación", level=2)
    doc.add_paragraph(
        "Si una falla atribuible a Andes Fibra deja a un cliente residencial sin servicio, la empresa debe restablecer "
        "el servicio en un máximo de 48 horas desde el reporte. El técnico puede llegar antes, pero ese es el plazo "
        "máximo comprometido."
    )
    doc.add_heading("Compensaciones", level=2)
    doc.add_paragraph(
        "El cliente residencial recibe un descuento equivalente a un día de su tarifa por cada día completo sin "
        "servicio, siempre que la caída supere las 6 horas continuas y sea atribuible a Andes Fibra."
    )
    doc.add_heading("Penalización por cancelación anticipada", level=2)
    doc.add_paragraph(
        "El contrato residencial dura 12 meses. Cancelar antes genera una penalización de 15 dólares por cada mes "
        "restante. La tabla muestra algunos ejemplos."
    )
    t = doc.add_table(rows=1, cols=3)
    t.style = "Table Grid"
    for celda, txt in zip(t.rows[0].cells, ("Mes de la cancelación", "Meses restantes", "Penalización (USD)")):
        celda.text = txt
    for mes, restan, pen in ((3, 9, 135), (6, 6, 90), (10, 2, 30)):
        fila = t.add_row().cells
        fila[0].text, fila[1].text, fila[2].text = str(mes), str(restan), str(pen)
    doc.add_heading("Uso aceptable", level=2)
    doc.add_paragraph(
        "El servicio residencial es para uso doméstico. Está prohibido revender el acceso a terceros o usarlo para "
        "operar un negocio con servidores públicos."
    )

    doc.add_heading("Parte B · Clientes Negocio", level=1)
    doc.add_heading("Tiempo de reparación", level=2)
    doc.add_paragraph(
        "Para los clientes del plan Negocio, Andes Fibra se compromete a una disponibilidad mensual del 99,5 % y a "
        "restablecer el servicio en un máximo de 4 horas desde el reporte, con técnico en sitio."
    )
    doc.add_heading("Compensaciones", level=2)
    doc.add_paragraph(
        "Si la caída supera las 4 horas, el cliente Negocio recibe un crédito del 10 % de su mensualidad por cada "
        "hora completa adicional, hasta un máximo del 100 % de la mensualidad."
    )
    doc.add_heading("Penalización por cancelación anticipada", level=2)
    doc.add_paragraph(
        "El contrato Negocio dura 24 meses. Cancelar antes genera una penalización de 25 dólares por cada mes "
        "restante."
    )
    doc.add_heading("Uso aceptable", level=2)
    doc.add_paragraph(
        "El plan Negocio incluye una dirección IP fija pública y permite alojar servidores propios, siempre que se "
        "respeten las leyes y no se afecte a otros clientes."
    )

    doc.add_heading("Protección de datos personales", level=1)
    doc.add_paragraph(
        "Andes Fibra conserva los datos de contacto y de facturación del cliente solo mientras dure el contrato y "
        "durante los 5 años posteriores exigidos por la ley fiscal imaginaria de este ejemplo."
    )
    doc.save(AQUI / "condiciones_servicio.docx")


def crear_pdf_con_texto() -> None:
    est = getSampleStyleSheet()
    pdf = SimpleDocTemplate(str(AQUI / "boletin_mantenimiento.pdf"), pagesize=A4,
                            leftMargin=2.2 * cm, rightMargin=2.2 * cm, topMargin=2 * cm, bottomMargin=2 * cm,
                            title="Boletín de mantenimiento programado", author="Andes Fibra (ficticio)")
    h, n = est["Heading1"], est["BodyText"]
    h2 = est["Heading2"]
    elementos = [
        Paragraph("Boletín de mantenimiento programado · noviembre de 2026", h),
        Paragraph("Documento ficticio creado para el curso. Andes Fibra realizará trabajos de ampliación en la red "
                  "de la central durante noviembre. Durante las ventanas indicadas, el servicio de internet y de "
                  "telefonía VoIP puede interrumpirse varias veces.", n),
        Spacer(1, 0.4 * cm),
        Paragraph("Ventanas por zona", h2),
    ]
    filas = [["Zona", "Fecha", "Horario"],
             ["Zona Norte (Los Pinos, Villa Real)", "4 de noviembre", "01:00 a 05:00"],
             ["Zona Sur (El Roble, San Gabriel)", "10 de noviembre", "00:00 a 04:00"],
             ["Zona Centro (Plaza Mayor)", "14 de noviembre", "02:00 a 06:00"]]
    tabla = Table(filas, hAlign="LEFT")
    tabla.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                               ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey)]))
    elementos += [tabla, Spacer(1, 0.4 * cm),
                  Paragraph("Estos trabajos se anuncian con 48 horas de anticipación. Una caída causada por "
                            "mantenimiento anunciado con esa anticipación no genera compensación. Durante la "
                            "ventana, la luz PON del módem óptico parpadeará mientras el equipo vuelve a "
                            "conectarse a la central.", n),
                  PageBreak(),
                  Paragraph("Preguntas frecuentes sobre el mantenimiento", h),
                  Paragraph("<b>¿Qué pasa con las llamadas de telefonía VoIP?</b> Las llamadas en curso se cortan al "
                            "iniciar la ventana y la línea queda sin tono hasta que termine el trabajo.", n),
                  Spacer(1, 0.3 * cm),
                  Paragraph("<b>¿Qué hago si el servicio no vuelve al terminar la ventana?</b> Espere 30 minutos "
                            "después del horario indicado, reinicie el router y, si el problema continúa, reporte el "
                            "caso al 0800-26337 indicando el código MNT-2611.", n),
                  Spacer(1, 0.3 * cm),
                  Paragraph("<b>¿Se pueden pedir cambios de fecha?</b> No. Las fechas son fijas porque afectan a "
                            "toda la zona. Los clientes del plan Negocio reciben además un aviso por correo con 5 "
                            "días de anticipación.", n)]
    pdf.build(elementos)


def crear_pdf_escaneado() -> None:
    """Un PDF que solo contiene una IMAGEN de texto (como un papel escaneado): no tiene capa de texto."""
    from PIL import Image, ImageDraw, ImageFont
    from reportlab.pdfgen import canvas

    img = Image.new("RGB", (1240, 400), "white")
    dib = ImageDraw.Draw(img)
    try:
        fuente = ImageFont.load_default(size=34)
    except TypeError:                       # Pillow antiguo
        fuente = ImageFont.load_default()
    lineas = ["ACTA DE INSTALACION (COPIA ESCANEADA)", "Cliente: Familia Rojas  Plan: Familiar 300 Mbps",
              "Equipo instalado: AX-2200   Metros de cable: 42", "Firma del tecnico: _______________"]
    for i, linea in enumerate(lineas):
        dib.text((40, 40 + i * 70), linea, fill="black", font=fuente)
    png = AQUI / "_escaneo_tmp.png"
    img.save(png)
    c = canvas.Canvas(str(AQUI / "acta_escaneada.pdf"), pagesize=A4)
    c.drawImage(str(png), 1.5 * cm, 18 * cm, width=18 * cm, height=18 * cm * 400 / 1240)
    c.save()
    png.unlink()


if __name__ == "__main__":
    crear_docx()
    crear_pdf_con_texto()
    crear_pdf_escaneado()
    print("Generados: condiciones_servicio.docx, boletin_mantenimiento.pdf, acta_escaneada.pdf")
